"""
Esportazione del listone in Excel.

Il servizio (`export_excel.listone_xlsx`) e' testato da solo, senza database:
costruisce un workbook da righe passate a mano. La route in piu' verifica solo
che il file scaricato sia coerente con quello che il servizio produce.
"""

import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.export_excel import INTESTAZIONI_LISTONE, listone_xlsx


def _righe(buffer_o_bytes):
    """Le righe del foglio (intestazione compresa), come tuple di valori."""
    sorgente = BytesIO(buffer_o_bytes) if isinstance(buffer_o_bytes, bytes) else buffer_o_bytes
    wb = load_workbook(sorgente)
    ws = wb.active
    return [tuple(cella.value for cella in riga) for riga in ws.iter_rows()]


class TestListoneXlsx:
    def test_intestazione(self):
        righe = _righe(listone_xlsx([]))
        assert righe == [tuple(INTESTAZIONI_LISTONE)]

    def test_una_riga_con_tutti_i_valori(self):
        giocatore = {
            "nome": "Osimhen",
            "squadra_att": "Napoli FC",
            "detentore_cartellino": "Napoli FC",
            "club": "NAP",
            "quot_att_mantra": 25,
            "tipo_contratto": "Indeterminato",
            "ruolo": "{A}",
            "costo": 120,
            "scadenza_contratto": datetime.date(2027, 6, 30),
            "data_nascita": datetime.date(1998, 12, 29),
            "valore_mercato": 75_000_000,
        }

        righe = _righe(listone_xlsx([giocatore]))

        assert righe[1] == (
            "Osimhen", "Napoli FC", "Napoli FC", "NAP", 25, "Indeterminato",
            "A", 120, datetime.datetime(2027, 6, 30), datetime.datetime(1998, 12, 29), 75.0,
        )

    def test_il_ruolo_grezzo_perde_le_graffe(self):
        """Il campo e' un array Postgres: '{DC,DD}' deve arrivare come 'DC,DD',
        non con le graffe che l'utente non deve vedere in un file scaricato."""
        giocatore = {"nome": "X", "squadra_att": None, "detentore_cartellino": None,
                     "club": None, "quot_att_mantra": None, "tipo_contratto": None,
                     "ruolo": "{DC,DD}", "costo": None, "scadenza_contratto": None,
                     "data_nascita": None, "valore_mercato": None}

        righe = _righe(listone_xlsx([giocatore]))

        assert righe[1][6] == "DC,DD"

    def test_il_valore_di_mercato_e_in_milioni_non_in_euro(self):
        giocatore = {"nome": "X", "squadra_att": None, "detentore_cartellino": None,
                     "club": None, "quot_att_mantra": None, "tipo_contratto": None,
                     "ruolo": None, "costo": None, "scadenza_contratto": None,
                     "data_nascita": None, "valore_mercato": 1_500_000}

        righe = _righe(listone_xlsx([giocatore]))

        assert righe[1][10] == 1.5

    def test_valore_di_mercato_non_sincronizzato_resta_vuoto(self):
        giocatore = {"nome": "X", "squadra_att": None, "detentore_cartellino": None,
                     "club": None, "quot_att_mantra": None, "tipo_contratto": None,
                     "ruolo": None, "costo": None, "scadenza_contratto": None,
                     "data_nascita": None, "valore_mercato": None}

        righe = _righe(listone_xlsx([giocatore]))

        assert righe[1][10] is None

    def test_una_riga_per_giocatore(self):
        giocatori = [
            {"nome": f"G{i}", "squadra_att": None, "detentore_cartellino": None,
             "club": None, "quot_att_mantra": None, "tipo_contratto": None,
             "ruolo": None, "costo": None, "scadenza_contratto": None,
             "data_nascita": None, "valore_mercato": None}
            for i in range(3)
        ]

        righe = _righe(listone_xlsx(giocatori))

        assert len(righe) == 1 + 3


class TestRouteExport:
    pytestmark = pytest.mark.db

    def test_scarica_un_file_excel(self, app):
        risposta = app.test_client().get("/listone/export")

        assert risposta.status_code == 200
        assert risposta.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "listone.xlsx" in risposta.headers["Content-Disposition"]

    def test_il_numero_di_righe_corrisponde_al_listone(self, app):
        pagina = app.test_client().get("/listone").get_data(as_text=True)
        import re
        totale = int(re.search(r'id="risultati-count">(\d+)<', pagina).group(1))

        righe = _righe(app.test_client().get("/listone/export").data)

        assert len(righe) - 1 == totale
