"""Esportazione di elenchi di giocatori in un file Excel.

Un solo formato oggi (il listone), ma isolato dal blueprint perche' costruire
il workbook non ha nulla a che fare con l'instradamento HTTP.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.domini.ruoli import pulisci_ruolo

INTESTAZIONI_LISTONE = [
    "Nome", "Squadra", "Detentore cartellino", "Club", "Quotazione",
    "Contratto", "Ruolo", "Costo", "Scadenza contratto", "Data di nascita",
    "Valore mercato (Mln)",
]


def _valore_mercato_mln(valore_euro):
    """Stesse unita' di misura mostrate a schermo (milioni), ma come numero:
    e' un file da ordinare e filtrare in Excel, non da leggere a video."""
    return round(valore_euro / 1_000_000, 2) if valore_euro is not None else None


def listone_xlsx(giocatori: list[dict]) -> BytesIO:
    """Un foglio con una riga per giocatore, colonne grezze del database.

    Niente stringhe pre-formattate ("27 anni (12/05/1998)", "12 Mln"): date e
    numeri restano tali, cosi' Excel li ordina e filtra invece di trattarli
    come testo.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Listone"

    ws.append(INTESTAZIONI_LISTONE)
    for cella in ws[1]:
        cella.font = Font(bold=True)

    for g in giocatori:
        ws.append([
            g["nome"],
            g["squadra_att"],
            g["detentore_cartellino"],
            g["club"],
            g["quot_att_mantra"],
            g["tipo_contratto"],
            pulisci_ruolo(g["ruolo"]),
            g["costo"],
            g["scadenza_contratto"],
            g["data_nascita"],
            _valore_mercato_mln(g["valore_mercato"]),
        ])

    for indice, intestazione in enumerate(INTESTAZIONI_LISTONE, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = max(12, len(intestazione) + 2)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
